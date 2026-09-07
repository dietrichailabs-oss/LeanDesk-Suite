from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook

from leandesk.sheets import SheetsFrame, WorkbookModel


def export(workbook, path: Path):
    frame = object.__new__(SheetsFrame)
    frame.workbook = workbook
    frame._save_xlsx(path)


@pytest.mark.parametrize("sheet_name", ["Sheet1", "r7", "Sales Data"])
def test_native_table_survives_actual_xlsx_export_edit_and_reexport(tmp_path, sheet_name):
    workbook = WorkbookModel()
    sheet = workbook.sheets[0]
    sheet.name = sheet_name
    sheet.cells.update({"A1": "Item", "B1": "Value", "A2": "Alpha",
                        "B2": "2", "A3": "Beta", "B3": "3"})
    workbook.office_features_for(sheet).create_table(sheet, "A1:B3", name="Table1")
    native = WorkbookModel.from_dict(workbook.to_dict())
    target = tmp_path / "table.xlsx"
    export(native, target)
    with ZipFile(target) as archive:
        assert archive.testzip() is None
        parts = [name for name in archive.namelist() if name.startswith("xl/tables/")]
        assert len(parts) == 1
        table_xml = ET.fromstring(archive.read(parts[0]))
        assert table_xml.attrib["ref"] == "A1:B3"
        assert table_xml.attrib["displayName"] == "Table1"
    reopened = SheetsFrame._load_xlsx(target)
    model = reopened.sheets[0]
    table = reopened.office_features_for(model).tables["Table1"]
    assert table.range_ref == "A1:B3"
    assert table.headers == ["Item", "Value"]
    assert table.style == "Medium2"
    assert table.banded_rows is True
    model.set("B2", "99")
    second = tmp_path / "edited.xlsx"
    export(reopened, second)
    external = load_workbook(second)
    try:
        assert external.active["B2"].value == "99"
        assert external.active.tables["Table1"].ref == "A1:B3"
    finally:
        external.close()

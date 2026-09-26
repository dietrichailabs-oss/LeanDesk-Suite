"""DOCX body tables must remain separate displayed blocks in Writer."""

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from leandesk.document_formats import LeanDocument
from leandesk.writer import WriterFrame


def body_blocks(path):
    document = Document(path)
    result = []
    for child in document.element.body:
        if child.tag == qn("w:p"):
            result.append(("p", Paragraph(child, document).text))
        elif child.tag == qn("w:tbl"):
            result.append(("t", Table(child, document).cell(0, 0).text))
    return result


@pytest.mark.parametrize(("blocks", "text", "indices"), [
    ([("p", "Before"), ("t", "A11"), ("p", "After")], "Before\n\nAfter", ["2.0"]),
    ([("t", "A11"), ("p", "After")], "\nAfter", ["1.0"]),
    ([("p", "Before"), ("t", "A11")], "Before\n", ["2.0"]),
    ([("t", "A11")], "", ["1.0"]),
    ([("t", "A11"), ("t", "B22"), ("p", "After")], "\n\nAfter", ["1.0", "2.0"]),
    ([("p", "Before"), ("t", "A11"), ("t", "B22"), ("p", "After")], "Before\n\n\nAfter", ["2.0", "3.0"]),
    ([("p", ""), ("t", "A11"), ("p", "")], "\n\n", ["2.0"]),
    ([("p", "Before"), ("p", ""), ("t", "A11"), ("p", ""), ("p", "After")], "Before\n\n\n\nAfter", ["3.0"]),
    ([("p", "\U0001f600 Before"), ("t", "A11"), ("p", "After")], "\U0001f600 Before\n\nAfter", ["2.0"]),
])
def test_import_tables_have_separate_lines_and_stable_reexport(tmp_path, blocks, text, indices):
    source = tmp_path / "original.docx"
    document = Document()
    for kind, value in blocks:
        if kind == "p":
            document.add_paragraph(value)
        else:
            document.add_table(rows=1, cols=1).cell(0, 0).text = value
    document.save(source)
    original_bytes = source.read_bytes()

    current = source
    for cycle in range(3):
        loaded = WriterFrame._load_docx(current)
        assert loaded.text == text
        assert [item["index"] for item in loaded.metadata["objects"]] == indices
        assert [item["data"][0][0] for item in loaded.metadata["objects"]] == [
            value for kind, value in blocks if kind == "t"
        ]
        current = tmp_path / f"cycle-{cycle}.docx"
        WriterFrame._save_docx(loaded, current)
        assert body_blocks(current) == blocks
    assert source.read_bytes() == original_bytes


def test_native_gui_table_fixture_keeps_paragraph_boundaries(tmp_path):
    before = "F2 Writer roundtrip 506ae762. Before table."
    after = "After table F2 end."
    cells = [["F2-A11", "F2-B22", ""], ["", "", ""], ["", "", ""]]
    native = LeanDocument(text=f"{before}\n\n{after}", metadata={"objects": [
        {"kind": "table", "rows": 3, "cols": 3, "data": cells, "index": "2.0"},
    ]})
    exported = tmp_path / "writer-table.docx"
    WriterFrame._save_docx(native, exported)
    assert body_blocks(exported) == [("p", before), ("t", "F2-A11"), ("p", after)]
    loaded = WriterFrame._load_docx(exported)
    assert loaded.text == native.text
    assert loaded.metadata["objects"][0]["index"] == "2.0"
    assert loaded.metadata["objects"][0]["data"] == cells


def test_heading_after_table_uses_its_own_correct_line(tmp_path):
    source = tmp_path / "headings.docx"
    document = Document()
    document.add_heading("Before", level=1)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "A11"
    document.add_heading("After", level=2)
    document.save(source)
    loaded = WriterFrame._load_docx(source)
    assert loaded.text == "Before\n\nAfter"
    assert [(tag.tag, tag.start, tag.end) for tag in loaded.tags] == [
        ("heading_1", "1.0", "1.end"), ("heading_2", "3.0", "3.end"),
    ]

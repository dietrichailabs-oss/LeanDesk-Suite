"""No accumulated export/import geometry drift in the canonical slide canvas."""
from types import SimpleNamespace

import pytest
from pptx import Presentation

from leandesk.slides import DeckModel, SlideModel, SlideObject, SlidesFrame


@pytest.mark.parametrize("x,width", [(80, 480), (12.345678, 480.12345)])
def test_repeated_actual_pptx_export_import_preserves_emu_geometry(tmp_path, x, width):
    item = SlideObject(kind="table", x=x, y=160, width=width, height=180,
                       data={"rows": 1, "cols": 2, "values": [["marker", "editable"]]})
    frame = SimpleNamespace(deck=DeckModel(slides=[SlideModel(title="Heading", body="Unchanged", objects=[item])]))
    expected = None
    for cycle in range(5):
        path = tmp_path / f"cycle-{cycle}.pptx"
        SlidesFrame._save_pptx(frame, path)
        ppt = Presentation(str(path))
        assert (ppt.slide_width, ppt.slide_height) == (12192000, 6858000)
        table = next(shape for shape in ppt.slides[0].shapes if shape.has_table)
        geometry = (table.left, table.top, table.width, table.height)
        if expected is None:
            expected = geometry
        assert geometry == expected
        assert table.table.cell(0, 0).text == "marker"
        assert len(ppt.slides[0].shapes) == 3
        frame.deck = SlidesFrame._load_pptx(path)
        assert frame.deck.slides[0].title == "Heading"
        assert frame.deck.slides[0].body == "Unchanged"

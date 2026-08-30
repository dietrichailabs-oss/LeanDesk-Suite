from __future__ import annotations

import io
import json
import os
from pathlib import Path
import shutil
import sqlite3
import stat
import tempfile
from datetime import datetime, timedelta, timezone
from unittest.mock import patch
import zipfile

import pytest

from leandesk.backup_integrity import BackupIntegrityError, safe_member, verify_backup_artifact
from leandesk.backup_service import create_backup, restore_backup
from leandesk.compatibility import cleanup_stale_conversion_roots
from leandesk.core import AppSettings, RecoveryRecord, RecoveryStore
from leandesk.data_boundary import DataBoundaryError, DataCorruptionError, UnsupportedSchemaVersion
from leandesk.document_formats import load_native
from leandesk.draw import Drawing, Shape
from leandesk.notes import Note
from leandesk.organizer import Task, _load_collection
from leandesk.save_policy import (
    ImportedSourceProtectionError,
    UnsupportedSaveFormatError,
    same_file_identity,
    validate_destination,
)
from leandesk.sheets import SheetModel, WorkbookModel, safe_number_expression
from leandesk.slides import DeckModel, SlideModel
from leandesk.updates.update_checker import check_for_updates, set_enabled
from leandesk.updates.update_manifest import MANIFEST_URL, ManifestError, parse_manifest
from leandesk.updates.version_compare import VersionError, compare_versions, is_newer

ROOT = Path(__file__).resolve().parents[1]


class FakeResponse:
    def __init__(self, body: bytes, *, final_url: str = MANIFEST_URL, status: int = 200):
        self._body = io.BytesIO(body)
        self._final_url = final_url
        self.status = status
        self.headers = {"Content-Length": str(len(body))}

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return False

    def read(self, size: int = -1) -> bytes:
        return self._body.read(size)

    def geturl(self) -> str:
        return self._final_url


def manifest_bytes(version: str = "0.8.1", **overrides) -> bytes:
    payload = {
        "product": "leandesk-suite",
        "latest_version": version,
        "release_name": f"LeanDesk Suite {version}",
        "published_at": "2026-08-23T00:00:00Z",
        "release_url": "https://www.dietrichailabs.com/apps/leandesk/",
        "download_url": "https://www.dietrichailabs.com/downloads/",
        "sha256": "",
        "message": "A new version of LeanDesk Suite is available.",
    }
    payload.update(overrides)
    return json.dumps(payload).encode("utf-8")


@pytest.mark.parametrize(
    ("candidate", "current", "expected"),
    [
        ("0.10.0", "0.9.0", True),
        ("0.8.1", "0.8.0", True),
        ("0.8.0", "0.8.0", False),
        ("0.7.9", "0.8.0", False),
        ("0.8.0-rc1", "0.8.0", False),
        ("0.8.0", "0.8.0-rc1", True),
        ("1.0.0-rc.10", "1.0.0-rc.2", True),
    ],
)
def test_semantic_version_order(candidate: str, current: str, expected: bool) -> None:
    assert is_newer(candidate, current) is expected


@pytest.mark.parametrize("value", ["", "1..2", "1.2-beta!", "01.2.3", "not-a-version"])
def test_malformed_versions_rejected(value: str) -> None:
    with pytest.raises(VersionError):
        compare_versions(value, "0.8.0")


def test_manifest_validates_exact_product_and_official_urls() -> None:
    parsed = parse_manifest(manifest_bytes("0.8.1", sha256="A" * 64))
    assert parsed.latest_version == "0.8.1"
    assert parsed.release_url == "https://www.dietrichailabs.com/apps/leandesk/"
    assert parsed.download_url == "https://www.dietrichailabs.com/downloads/"
    assert parsed.sha256 == "A" * 64


@pytest.mark.parametrize(
    "changes",
    [
        {"product": "other"},
        {"latest_version": None},
        {"latest_version": "bad version"},
        {"release_url": "https://evil.example/apps/leandesk/"},
        {"release_url": "https://www.dietrichailabs.com/apps/leandesk-evil/"},
        {"download_url": "http://www.dietrichailabs.com/downloads/"},
        {"download_url": "https://www.dietrichailabs.com/downloadsville/"},
        {"release_url": "", "download_url": ""},
        {"sha256": "not-a-hash"},
        {"published_at": "yesterday"},
    ],
)
def test_manifest_rejects_invalid_fields(changes: dict) -> None:
    with pytest.raises(ManifestError):
        parse_manifest(manifest_bytes(**changes))


def test_first_automatic_update_check_runs(tmp_path: Path) -> None:
    calls = []
    result = check_for_updates(
        "0.8.0",
        opener=lambda req, timeout: calls.append((req.full_url, timeout)) or FakeResponse(manifest_bytes()),
        state_path=tmp_path / "state.json",
        now=datetime(2026, 8, 23, tzinfo=timezone.utc),
    )
    assert result.status == "update_available"
    assert result.checked is True
    assert calls == [(MANIFEST_URL, 5.0)]


def test_no_second_automatic_check_within_seven_days(tmp_path: Path) -> None:
    calls = []
    opener = lambda req, timeout: calls.append(req.full_url) or FakeResponse(manifest_bytes())
    now = datetime(2026, 8, 23, tzinfo=timezone.utc)
    check_for_updates("0.8.0", opener=opener, state_path=tmp_path / "state.json", now=now)
    result = check_for_updates("0.8.0", opener=opener, state_path=tmp_path / "state.json", now=now + timedelta(days=6, hours=23))
    assert result.status == "not_due"
    assert len(calls) == 1


def test_automatic_check_after_seven_days(tmp_path: Path) -> None:
    calls = []
    opener = lambda req, timeout: calls.append(req.full_url) or FakeResponse(manifest_bytes())
    now = datetime(2026, 8, 23, tzinfo=timezone.utc)
    check_for_updates("0.8.0", opener=opener, state_path=tmp_path / "state.json", now=now)
    result = check_for_updates("0.8.0", opener=opener, state_path=tmp_path / "state.json", now=now + timedelta(days=7))
    assert result.status == "update_available"
    assert len(calls) == 2


def test_manual_check_bypasses_timer_and_disabled_preference(tmp_path: Path) -> None:
    state = tmp_path / "state.json"
    set_enabled(False, state_path=state)
    calls = []
    result = check_for_updates(
        "0.8.0",
        force=True,
        opener=lambda req, timeout: calls.append(req.full_url) or FakeResponse(manifest_bytes("0.8.0")),
        state_path=state,
        now=datetime(2026, 8, 23, tzinfo=timezone.utc),
    )
    assert result.status == "current"
    assert calls == [MANIFEST_URL]


def test_disabled_automatic_check_makes_no_request(tmp_path: Path) -> None:
    state = tmp_path / "state.json"
    set_enabled(False, state_path=state)
    calls = []
    result = check_for_updates("0.8.0", opener=lambda *_a, **_k: calls.append(1), state_path=state)
    assert result.status == "disabled"
    assert not calls


def test_failed_attempt_is_rate_limited_to_prevent_offline_polling(tmp_path: Path) -> None:
    state = tmp_path / "state.json"
    now = datetime(2026, 8, 23, tzinfo=timezone.utc)
    calls = []

    def fail(*_args, **_kwargs):
        calls.append(1)
        raise OSError("offline")

    first = check_for_updates("0.8.0", opener=fail, state_path=state, now=now)
    second = check_for_updates("0.8.0", opener=fail, state_path=state, now=now + timedelta(hours=1))
    assert first.status == "error"
    assert second.status == "not_due"
    assert len(calls) == 1


def test_update_checker_rejects_nonofficial_final_redirect(tmp_path: Path) -> None:
    result = check_for_updates(
        "0.8.0",
        opener=lambda *_a, **_k: FakeResponse(manifest_bytes(), final_url="https://evil.example/manifest.json"),
        state_path=tmp_path / "state.json",
    )
    assert result.status == "error"


@pytest.mark.parametrize(
    "body",
    [b"{bad", json.dumps({"product": "leandesk-suite"}).encode(), manifest_bytes(latest_version="bad")],
)
def test_update_checker_fails_closed_on_invalid_manifest(tmp_path: Path, body: bytes) -> None:
    result = check_for_updates(
        "0.8.0",
        opener=lambda *_a, **_k: FakeResponse(body),
        state_path=tmp_path / "state.json",
    )
    assert result.status == "error"
    assert result.checked is True


def test_update_checker_handles_http_error_status(tmp_path: Path) -> None:
    result = check_for_updates(
        "0.8.0",
        opener=lambda *_a, **_k: FakeResponse(manifest_bytes(), status=503),
        state_path=tmp_path / "state.json",
    )
    assert result.status == "error"


@pytest.mark.parametrize("server_version", ["0.8.0", "0.7.9"])
def test_update_checker_ignores_equal_or_older_server_versions(tmp_path: Path, server_version: str) -> None:
    result = check_for_updates(
        "0.8.0",
        opener=lambda *_a, **_k: FakeResponse(manifest_bytes(server_version)),
        state_path=tmp_path / "state.json",
    )
    assert result.status == "current"
    assert result.latest_version == server_version


def test_malformed_update_state_is_preserved_and_prevents_network(tmp_path: Path) -> None:
    state = tmp_path / "state.json"
    original = b"{malformed update state"
    state.write_bytes(original)
    calls = []
    result = check_for_updates(
        "0.8.0",
        opener=lambda *_a, **_k: calls.append(1) or FakeResponse(manifest_bytes()),
        state_path=state,
    )
    assert result.status == "error"
    assert not calls
    assert state.read_bytes() == original


def test_save_policy_rejects_exact_imported_source(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    source.write_bytes(b"source")
    with pytest.raises(ImportedSourceProtectionError):
        validate_destination("Writer", source, imported_source=source)
    assert source.read_bytes() == b"source"


def test_save_policy_rejects_hardlink_alias(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    alias = tmp_path / "alias.docx"
    source.write_bytes(b"source")
    os.link(source, alias)
    assert same_file_identity(source, alias)
    with pytest.raises(ImportedSourceProtectionError):
        validate_destination("Writer", alias, imported_source=source)


def test_save_policy_rejects_symlink_alias_when_supported(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    alias = tmp_path / "alias.xlsx"
    source.write_bytes(b"source")
    try:
        alias.symlink_to(source)
    except (OSError, NotImplementedError):
        pytest.skip("symlinks unavailable")
    assert same_file_identity(source, alias)
    with pytest.raises(ImportedSourceProtectionError):
        validate_destination("Sheets", alias, imported_source=source)


@pytest.mark.parametrize(
    ("module", "suffix"),
    [("Writer", ".odt"), ("Writer", ".doc"), ("Sheets", ".ods"), ("Sheets", ".xls"), ("Slides", ".odp"), ("Slides", ".ppt")],
)
def test_import_only_suffixes_are_not_writable(module: str, suffix: str, tmp_path: Path) -> None:
    with pytest.raises(UnsupportedSaveFormatError):
        validate_destination(module, tmp_path / f"out{suffix}")


@pytest.mark.parametrize("identifier", ["../outside", "..", ".", "a/b", r"a\\b", "C:escape", "CON", "name.json", "bad\x00id"])
def test_recovery_identifiers_reject_traversal_and_windows_ambiguity(identifier: str, tmp_path: Path) -> None:
    store = RecoveryStore(tmp_path / "Recovery")
    with pytest.raises(ValueError):
        store.delete(identifier)


@pytest.mark.parametrize("module", ["Writer", "Sheets", "Slides", "Notes", "Draw", "Tasks", "Calendar", "Contacts"])
def test_recovery_store_supports_all_advertised_modules(module: str, tmp_path: Path) -> None:
    store = RecoveryStore(tmp_path / "Recovery")
    record = RecoveryRecord(f"id_{module.lower()}", module, module, "", "2026-08-23T00:00:00", {"module": module})
    store.save(record)
    rows = store.list(module)
    assert len(rows) == 1 and rows[0].module == module


def test_damaged_recovery_record_is_quarantined(tmp_path: Path) -> None:
    root = tmp_path / "Recovery"
    root.mkdir()
    damaged = root / "broken.json"
    damaged.write_text("{broken", encoding="utf-8")
    store = RecoveryStore(root)
    assert store.list() == []
    assert not damaged.exists()
    assert any((root / "Quarantine").iterdir())


def test_future_settings_remain_read_only_and_byte_identical(tmp_path: Path) -> None:
    path = tmp_path / "settings.json"
    original = b'{"schema_version":999,"future_field":{"keep":true},"theme":"Future"}'
    path.write_bytes(original)
    settings = AppSettings.load(path)
    assert settings.read_only
    settings.theme = "Changed"
    assert settings.save(path) is False
    assert path.read_bytes() == original


def test_malformed_settings_types_are_controlled_and_non_destructive(tmp_path: Path) -> None:
    path = tmp_path / "settings.json"
    original = b'{"schema_version":1,"autosave_seconds":"not-a-number","auto_check_updates":"yes"}'
    path.write_bytes(original)
    settings = AppSettings.load(path)
    assert settings.read_only
    assert settings.autosave_seconds == 30
    assert settings.auto_check_updates is True
    assert settings.save(path) is False
    assert path.read_bytes() == original


def test_recovery_record_with_non_object_payload_is_quarantined(tmp_path: Path) -> None:
    root = tmp_path / "Recovery"
    root.mkdir()
    path = root / "broken.json"
    path.write_text(
        json.dumps({
            "schema_version": 1,
            "recovery_id": "broken",
            "module": "Writer",
            "title": "x",
            "original_path": "",
            "saved_at": "2026-08-23T00:00:00Z",
            "payload": [],
        }),
        encoding="utf-8",
    )
    store = RecoveryStore(root)
    assert store.list() == []
    assert not path.exists()
    assert any((root / "Quarantine").iterdir())


def test_note_row_wrong_types_are_rejected_instead_of_coerced() -> None:
    with pytest.raises(DataCorruptionError):
        Note.from_dict({
            "note_id": str(__import__("uuid").uuid4()),
            "title": 123,
            "body": "x",
            "notebook": "General",
            "tags": "",
            "created_at": "",
            "updated_at": "",
            "pinned": "yes",
        })


def test_organizer_wrong_types_are_read_only_and_non_destructive(tmp_path: Path) -> None:
    path = tmp_path / "tasks.json"
    original = json.dumps({
        "schema_version": 1,
        "tasks": [{
            "task_id": str(__import__("uuid").uuid4()),
            "title": 123,
            "due": "",
            "priority": "Normal",
            "status": "Open",
            "project": "",
            "notes": "",
        }],
    }).encode("utf-8")
    path.write_bytes(original)
    items, read_only, error, _extra = _load_collection(path, "tasks", Task)
    assert items == []
    assert read_only is True
    assert error
    assert path.read_bytes() == original


@pytest.mark.parametrize(
    "payload,expected_exception",
    [
        (b'{"format_version":1,"title":"a","title":"b","text":"x","tags":[],"metadata":{}}', DataBoundaryError),
        (b'{"format_version":999,"title":"a","text":"x","tags":[],"metadata":{}}', UnsupportedSchemaVersion),
    ],
)
def test_native_document_corruption_or_future_version_is_non_destructive(tmp_path: Path, payload: bytes, expected_exception: type[Exception]) -> None:
    path = tmp_path / "doc.ldoc"
    path.write_bytes(payload)
    with pytest.raises(expected_exception):
        load_native(path)
    assert path.read_bytes() == payload


@pytest.mark.parametrize(
    "payload",
    [
        {"format_version": 1, "title": 99, "sheets": []},
        {"format_version": 1, "title": "x", "sheets": ["not-an-object"]},
        {"format_version": 1, "title": "x", "sheets": [{"name": "Sheet1", "cells": {"A1": 7}, "column_widths": {}}]},
        {"format_version": 1, "title": "x", "sheets": [{"name": "Sheet1", "cells": {}, "column_widths": {"A": "100"}}]},
    ],
)
def test_workbook_wrong_types_are_rejected_instead_of_coerced(payload: dict) -> None:
    with pytest.raises(DataCorruptionError):
        WorkbookModel.from_dict(payload)


@pytest.mark.parametrize(
    "payload",
    [
        {"format_version": 1, "title": 99, "slides": []},
        {"format_version": 1, "title": "x", "slides": ["not-an-object"]},
        {"format_version": 1, "title": "x", "slides": [{"title": 7, "body": "x"}]},
        {"format_version": 1, "title": "x", "slides": [{"title": "x", "body": "y", "image_media_type": "image/png"}]},
    ],
)
def test_deck_wrong_types_are_rejected_instead_of_coerced(payload: dict) -> None:
    with pytest.raises(DataCorruptionError):
        DeckModel.from_dict(payload)


@pytest.mark.parametrize(
    "payload",
    [
        {"format_version": 1, "title": 99, "width": 1200, "height": 760, "background": "#fff", "shapes": []},
        {"format_version": 1, "title": "x", "width": "1200", "height": 760, "background": "#fff", "shapes": []},
        {"format_version": 1, "title": "x", "width": 1200, "height": 760, "background": "#fff", "shapes": [{"shape_id": "1", "kind": "line", "x1": True, "y1": 0, "x2": 1, "y2": 1}]},
        {"format_version": 1, "title": "x", "width": 1200, "height": 760, "background": "#fff", "shapes": [{"shape_id": "1", "kind": "text", "x1": 0, "y1": 0, "x2": 1, "y2": 1, "text": 77}]},
    ],
)
def test_drawing_wrong_types_are_rejected_instead_of_coerced(payload: dict) -> None:
    with pytest.raises(DataCorruptionError):
        Drawing.from_dict(payload)


def test_formula_exponent_is_bounded_without_overflow() -> None:
    with pytest.raises(ValueError):
        safe_number_expression("2**1000000")


def test_formula_complexity_is_bounded() -> None:
    with pytest.raises(ValueError):
        safe_number_expression("+".join("1" for _ in range(300)))


def test_spreadsheet_grid_accepts_aa101_and_rejects_hidden_beyond_bounds() -> None:
    sheet = SheetModel()
    sheet.set("AA101", "visible")
    assert sheet.raw("AA101") == "visible"
    with pytest.raises(ValueError):
        sheet.set("BA1", "hidden")
    with pytest.raises(ValueError):
        sheet.set("A201", "hidden")


def test_writer_docx_loader_accepts_in_memory_compatibility_payload() -> None:
    from docx import Document
    from leandesk.writer import WriterFrame

    payload = io.BytesIO()
    document = Document()
    document.add_paragraph("Converted legacy content")
    document.save(payload)
    payload.seek(0)

    loaded = WriterFrame._load_docx(payload, title="Legacy Source")
    assert loaded.title == "Legacy Source"
    assert "Converted legacy content" in loaded.text


def test_slide_loader_discards_external_image_path() -> None:
    slide = SlideModel.from_dict({"title": "x", "body": "y", "image_path": "/etc/passwd"})
    assert slide.image_path == ""
    assert "image_path" not in slide.to_dict()


def test_stale_compatibility_roots_are_cleaned(tmp_path: Path) -> None:
    stale = tmp_path / "leandesk_compat_stale"
    stale.mkdir()
    (stale / "document.docx").write_bytes(b"sensitive")
    old = datetime.now().timestamp() - 3 * 24 * 3600
    os.utime(stale, (old, old))
    assert cleanup_stale_conversion_roots(tmp_path, older_than_seconds=3600) == 1
    assert not stale.exists()


@pytest.mark.parametrize("name", ["C:/escape.json", r"C:\\escape.json", "../escape.json", "/escape.json", r"\\server\\share\\file.json", "folder/CON.json"])
def test_backup_member_path_rejects_windows_and_traversal_forms(name: str) -> None:
    with pytest.raises(BackupIntegrityError):
        safe_member(name)


def test_backup_rejects_corrupt_sqlite_with_controlled_error(tmp_path: Path) -> None:
    archive = tmp_path / "bad.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("notes.db", b"not sqlite")
    with pytest.raises(BackupIntegrityError):
        verify_backup_artifact(archive)


def test_backup_rejects_duplicate_nested_native_members(tmp_path: Path) -> None:
    nested = io.BytesIO()
    with zipfile.ZipFile(nested, "w") as zf:
        zf.writestr("document.json", "{}")
        zf.writestr("document.json", '{"other":1}')
    archive = tmp_path / "bad.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("doc.ldoc", nested.getvalue())
    with pytest.raises(BackupIntegrityError):
        verify_backup_artifact(archive)


def test_backup_directory_rejects_linked_directory(tmp_path: Path) -> None:
    root = tmp_path / "profile"
    real = tmp_path / "real"
    root.mkdir(); real.mkdir()
    (real / "settings.json").write_text("{}", encoding="utf-8")
    try:
        (root / "linked").symlink_to(real, target_is_directory=True)
    except (OSError, NotImplementedError):
        pytest.skip("symlinks unavailable")
    with pytest.raises(BackupIntegrityError):
        verify_backup_artifact(root)


def test_backup_create_and_staged_restore_round_trip(tmp_path: Path) -> None:
    source = tmp_path / "profile"
    source.mkdir()
    (source / "settings.json").write_text('{"schema_version":1,"theme":"Midnight"}', encoding="utf-8")
    (source / "notes.json").write_text('{"schema_version":1,"notes":[]}', encoding="utf-8")
    archive = tmp_path / "profile.ldbackup"
    created = create_backup(archive, data_root=source)
    assert created["valid"] and created["files"] == 2
    restored = tmp_path / "restored"
    result = restore_backup(archive, data_root=restored)
    assert result["valid"]
    assert (restored / "settings.json").read_text(encoding="utf-8") == (source / "settings.json").read_text(encoding="utf-8")
    assert (restored / "notes.json").is_file()


def test_failed_restore_leaves_existing_profile_unchanged(tmp_path: Path) -> None:
    target = tmp_path / "profile"
    target.mkdir()
    sentinel = target / "settings.json"
    sentinel.write_text('{"keep":true}', encoding="utf-8")
    invalid = tmp_path / "invalid.zip"
    with zipfile.ZipFile(invalid, "w") as zf:
        zf.writestr("../escape.json", "{}")
    before = sentinel.read_bytes()
    with pytest.raises(BackupIntegrityError):
        restore_backup(invalid, data_root=target)
    assert sentinel.read_bytes() == before


def test_update_notification_is_restrained_to_one_automatic_notice_per_version() -> None:
    from leandesk.app import LeanDeskApp
    from leandesk.updates.update_checker import UpdateResult

    app = object.__new__(LeanDeskApp)
    app._notified_update_versions = set()
    shown: list[str] = []
    app._show_update_available = lambda result: shown.append(result.latest_version or "")

    available = UpdateResult(
        "update_available",
        True,
        "0.8.0",
        latest_version="0.8.1",
        release_url="https://www.dietrichailabs.com/apps/leandesk/",
    )
    app._handle_update_result(available, manual=False)
    app._handle_update_result(available, manual=False)
    app._handle_update_result(UpdateResult("current", True, "0.8.0", latest_version="0.8.0"), manual=False)
    app._handle_update_result(UpdateResult("error", True, "0.8.0", error="offline"), manual=False)
    assert shown == ["0.8.1"]


def test_update_feature_is_wired_into_app_and_contains_no_downloader() -> None:
    app_source = (ROOT / "leandesk" / "app.py").read_text(encoding="utf-8")
    checker_source = (ROOT / "leandesk" / "updates" / "update_checker.py").read_text(encoding="utf-8")
    assert "self.after(1500, self._schedule_automatic_update_check)" in app_source
    assert "cleanup_stale_conversion_roots()" in app_source
    assert "Automatically check for LeanDesk updates once a week" in app_source
    assert "Check for Updates Now" in app_source
    manifest_source = (ROOT / "leandesk" / "updates" / "update_manifest.py").read_text(encoding="utf-8")
    assert "from .update_checker import MANIFEST_URL" in app_source
    assert MANIFEST_URL in manifest_source
    forbidden = ("subprocess", "Popen(", "os.startfile", "urlretrieve", "shutil.copyfileobj")
    assert not any(token in checker_source for token in forbidden)


def test_export_commands_use_guarded_atomic_write_boundaries() -> None:
    writer = (ROOT / "leandesk" / "writer.py").read_text(encoding="utf-8")
    sheets = (ROOT / "leandesk" / "sheets.py").read_text(encoding="utf-8")
    slides = (ROOT / "leandesk" / "slides.py").read_text(encoding="utf-8")
    assert 'return self._write_to(Path(value))' in writer
    assert 'return bool(value) and self._write(Path(value))' in sheets
    assert 'return bool(value) and self._write(Path(value))' in slides
    assert 'self._save_csv(Path(value))' not in sheets
    assert 'self._save_pptx(Path(value))' not in slides


def test_dependency_lock_is_exact_and_build_consumes_it() -> None:
    lock = [line.strip() for line in (ROOT / "requirements.lock.txt").read_text(encoding="utf-8").splitlines() if line.strip()]
    assert lock
    assert all("==" in line and not any(op in line for op in (">=", "<=", "~=", "<", ">")) for line in lock)
    build = (ROOT / "BUILD_LEANDESK_SUITE.ps1").read_text(encoding="utf-8-sig")
    assert "requirements.lock.txt" in build
    assert "tools\\run_authoritative_tests.py" in build
    assert "-m unittest -v test_leandesk.py" not in build


def test_canonical_runner_is_recursive_and_enforces_required_suites() -> None:
    source = (ROOT / "tools" / "run_authoritative_tests.py").read_text(encoding="utf-8")
    assert '"test_leandesk.py"' in source
    assert '"test_compatibility.py"' in source
    assert '"tests/test_correction_2_qa.py"' in source
    assert '"--collect-only", "-q"' in source
    assert '[sys.executable, "-m", "pytest", "-q"]' in source
    assert 'str(root / "tests")' not in source


@pytest.mark.skipif(not os.environ.get("DISPLAY") and os.name != "nt", reason="GUI test requires an active display/Xvfb")
@pytest.mark.parametrize("module", ["Writer", "Sheets", "Slides"])
def test_actual_gui_save_as_same_imported_source_is_controlled_and_non_destructive(module: str, tmp_path: Path) -> None:
    import tkinter as tk
    from leandesk.core import RecentFiles
    from leandesk.writer import WriterFrame
    from leandesk.sheets import SheetsFrame
    from leandesk.slides import SlidesFrame

    root = tk.Tk(); root.withdraw()
    recent = RecentFiles(tmp_path / "recent.json")
    settings = AppSettings()
    try:
        if module == "Writer":
            frame = WriterFrame(root, recent=recent, settings=settings)
            source = tmp_path / "foreign.docx"
            source.write_bytes(b"immutable foreign bytes")
            frame.current_path = source; frame.imported_source_path = source; frame.dirty = True
            target = "leandesk.writer.filedialog.asksaveasfilename"
        elif module == "Sheets":
            frame = SheetsFrame(root, recent=recent)
            source = tmp_path / "foreign.xlsx"
            source.write_bytes(b"immutable foreign bytes")
            frame.current_path = source; frame.imported_source_path = source; frame.dirty = True
            target = "leandesk.sheets.filedialog.asksaveasfilename"
        else:
            frame = SlidesFrame(root, recent=recent)
            source = tmp_path / "foreign.pptx"
            source.write_bytes(b"immutable foreign bytes")
            frame.current_path = source; frame.imported_source_path = source; frame.dirty = True
            target = "leandesk.slides.filedialog.asksaveasfilename"
        frame.recovery = RecoveryStore(tmp_path / "Recovery")
        before = source.read_bytes()
        with patch(target, return_value=str(source)), patch("tkinter.messagebox.showwarning", return_value=None):
            assert frame.save_as() is False
        assert source.read_bytes() == before
    finally:
        root.destroy()

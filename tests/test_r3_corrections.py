"""R3 reproductions. Source tests are not signed Windows lifecycle evidence."""
import json
import os
from pathlib import Path
import subprocess
import sys
from types import SimpleNamespace

import pytest

from leandesk.slides import DeckModel, SlideModel, SlideObject, SlidesFrame
from leandesk.themes import SUITE_THEMES


THEME_PROCESS = r'''
import json, sys
from tkinter import ttk
from leandesk.app import LeanDeskApp
from leandesk.core import AppSettings, SETTINGS_FILE
from leandesk.ui import COLORS

def descendants(widget):
    for child in widget.winfo_children():
        yield child
        yield from descendants(child)

settings = AppSettings.load()
settings.auto_check_updates = False
settings.save()
app = LeanDeskApp()
app.show_settings()
app.update()
picker = next(w for w in descendants(app.content)
              if isinstance(w, ttk.Combobox) and 'Lavender Office' in w.cget('values'))
if sys.argv[1] == 'choose':
    picker.set(sys.argv[2])
    picker.event_generate('<<ComboboxSelected>>')
    app.update()
snapshot = {'selector': picker.get(), 'memory': app.settings.theme,
            'palette': dict(COLORS), 'background': app.cget('background')}
app.on_close()
snapshot['disk'] = json.loads(SETTINGS_FILE.read_text(encoding='utf-8'))['theme']
print('R3_THEME_RESULT=' + json.dumps(snapshot, sort_keys=True))
'''


@pytest.mark.parametrize('theme', tuple(SUITE_THEMES))
def test_theme_selection_survives_actual_close_and_new_process(tmp_path, theme):
    environment = os.environ.copy()
    environment['LOCALAPPDATA'] = str(tmp_path / 'profile')
    environment['LEANDESK_GUI_REPRO_PROFILE'] = environment['LOCALAPPDATA']
    environment['PYTHONDONTWRITEBYTECODE'] = '1'
    def run(mode):
        result = subprocess.run([sys.executable, '-c', THEME_PROCESS, mode, theme],
                                env=environment, capture_output=True, text=True, timeout=90)
        assert result.returncode == 0, result.stdout + result.stderr
        lines = [line for line in result.stdout.splitlines() if line.startswith('R3_THEME_RESULT=')]
        assert len(lines) == 1, result.stdout + result.stderr
        return json.loads(lines[0].split('=', 1)[1])
    chosen, reopened = run('choose'), run('reopen')
    assert chosen['selector'] == chosen['memory'] == chosen['disk'] == theme
    assert reopened == chosen


@pytest.mark.parametrize('rows,cols', [(2, 3), (3, 2), (4, 4)])
def test_pptx_editable_table_survives_native_export_import_and_second_export(tmp_path, rows, cols):
    values = [[f'R{row + 1}C{col + 1}-unique' for col in range(cols)] for row in range(rows)]
    model = SlideModel(title='R3 title sentinel', body='R3 body sentinel', notes='R3 notes sentinel')
    model.objects = [
        SlideObject(kind='text', x=50, y=130, width=200, height=50, text='before table'),
        SlideObject(kind='table', x=80, y=230, width=420, height=180,
                    data={'rows': rows, 'cols': cols, 'values': values}),
        SlideObject(kind='shape', x=570, y=280, width=180, height=100, text='after table'),
    ]
    frame = object.__new__(SlidesFrame)
    frame.deck = DeckModel.from_dict(DeckModel('R3 table fixture', [model]).to_dict())
    for iteration in range(2):
        path = tmp_path / f'table-roundtrip-{iteration}.pptx'
        frame._save_pptx(path)
        from pptx import Presentation
        exported = Presentation(path)
        table = next(shape.table for shape in exported.slides[0].shapes if shape.has_table)
        assert [[cell.text for cell in row.cells] for row in table.rows] == values
        frame.deck = SlidesFrame._load_pptx(path)
        slide = frame.deck.slides[0]
        assert slide.title == model.title and slide.body == model.body
        assert model.notes in slide.notes
        assert [item.kind for item in slide.objects] == ['text', 'table', 'shape']
        restored = slide.objects[1]
        assert restored.data == {'rows': rows, 'cols': cols, 'values': values}
        assert (restored.x, restored.y, restored.width, restored.height) == pytest.approx((80, 230, 420, 180), abs=.1)
        assert slide.objects[0].text == 'before table'
        assert slide.objects[2].text == 'after table'
        # Editability is a model operation, not a rendered screenshot assertion.
        values[0][0] = f'edited-{iteration}'
        restored.data['values'][0][0] = values[0][0]


@pytest.mark.parametrize('stdout', ['[]', 'null', '"not an object"', '{broken'])
def test_print_invalid_helper_response_is_controlled_and_cleans_temp(monkeypatch, stdout):
    from leandesk.document_formats import LeanDocument
    from leandesk.windows_print import PrintUnavailableError, print_rtf_document
    paths = []
    def run(command, **kwargs):
        path = Path(kwargs['env']['LEANDESK_PRINT_RTF'])
        assert path.is_file()
        paths.append(path)
        return SimpleNamespace(returncode=0, stdout=stdout, stderr='')
    monkeypatch.setattr(subprocess, 'run', run)
    with pytest.raises(PrintUnavailableError):
        print_rtf_document(LeanDocument())
    assert paths and all(not path.parent.exists() for path in paths)

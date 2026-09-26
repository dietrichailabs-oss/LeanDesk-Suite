"""Preserve the actual Writer text/table sequence across DOCX boundaries."""

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from leandesk.document_formats import LeanDocument
from leandesk.writer import WriterFrame


def table(index: str, value: str = "cell") -> dict:
    return {"kind": "table", "rows": 1, "cols": 1, "data": [[value]], "index": index}


def blocks(path: Path) -> list[tuple[str, str]]:
    document = Document(path)
    result = []
    for child in document.element.body:
        if child.tag == qn("w:p"):
            result.append(("text", Paragraph(child, document).text))
        elif child.tag == qn("w:tbl"):
            result.append(("table", Table(child, document).cell(0, 0).text))
    return result


@pytest.mark.parametrize(("text", "objects", "expected"), [
    ("*", [table("1.0")], [("table", "cell"), ("text", "*")]),
    ("*", [table("1.1")], [("text", "*"), ("table", "cell")]),
    ("beforeafter", [table("1.6")], [("text", "before"), ("table", "cell"), ("text", "after")]),
    ("*", [table("1.0", "one"), table("1.1", "two")], [("table", "one"), ("table", "two"), ("text", "*")]),
    ("*", [table("1.2", "two"), table("1.0", "one")], [("table", "one"), ("text", "*"), ("table", "two")]),
    ("\U0001f600after", [table("1.2")], [("text", "\U0001f600"), ("table", "cell"), ("text", "after")]),
    ("before\nafter", [table("end-1c")], [("text", "before"), ("text", "after"), ("table", "cell")]),
])
def test_docx_export_keeps_native_object_order(tmp_path, text, objects, expected):
    path = tmp_path / "ordered.docx"
    WriterFrame._save_docx(LeanDocument(text=text, metadata={"objects": objects}), path)
    assert blocks(path) == expected


@pytest.mark.parametrize("expected", [
    [("table", "first"), ("text", "after")],
    [("text", "before"), ("table", "last")],
    [("text", "before"), ("table", "middle"), ("text", "after")],
    [("table", "one"), ("table", "two"), ("text", "after")],
    [("text", "\U0001f600"), ("table", "unicode"), ("text", "after")],
])
def test_docx_import_and_reexport_keep_body_order(tmp_path, expected):
    original = tmp_path / "original.docx"
    document = Document()
    for kind, value in expected:
        if kind == "text":
            document.add_paragraph(value)
        else:
            document.add_table(rows=1, cols=1).cell(0, 0).text = value
    document.save(original)
    before = original.read_bytes()
    imported = WriterFrame._load_docx(original)
    assert [item["data"][0][0] for item in imported.metadata["objects"]] == [
        value for kind, value in expected if kind == "table"
    ]
    exported = tmp_path / "reexported.docx"
    WriterFrame._save_docx(imported, exported)
    assert blocks(exported) == expected
    assert original.read_bytes() == before
